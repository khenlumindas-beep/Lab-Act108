# -*- coding: utf-8 -*-
"""
Generates the filled-out IT108 Activity 7 laboratory report (DOCX) for the
SmartTrack system.

Usage:
    backend\\venv\\Scripts\\python.exe docs\\tools\\generate_activity7_report.py

Output:
    docs/IT108_Activity_7_SmartTrack_COMPLETED.docx

All content is taken from the actual project artifacts:
    - database/security/01_roles_and_grants.sql   (Part C)
    - database/security/02_rls_policy.sql         (Part D)
    - database/security/evidence/privilege_test_results.txt (test results)
    - backend/config.py, backend/.env.example, .gitignore   (Part E)
    - backend/routes/auth.py                      (Part F)
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "docs" / "IT108_Activity_7_SmartTrack_COMPLETED.docx"

SYSTEM_TITLE = ("SmartTrack — AI-Assisted International Mobility & Foreign "
                "Travel Management System (CSUCC Foreign Mobility Monitoring System)")

doc = Document()


# ---------------------------------------------------------------- helpers --
def title(text, size=14):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    return par


def center(text, size=11, bold=False, italic=False):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return par


def h(text, level=1):
    return doc.add_heading(text, level=level)


def p(text, bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


def sql(text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    return par


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
    doc.add_paragraph()
    return t


# ----------------------------------------------------------------- header --
title("COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY", 12)
title("CARAGA STATE UNIVERSITY — CABADBARAN CAMPUS", 12)
center("IT 108 — Advanced Database Management System", 11, bold=True)
center("Laboratory Activity No. 7 — PostgreSQL Database Security and Access Control", 11, bold=True)
center("Roles, Privileges, Row-Level Security, Credential Protection, and SQL Injection Prevention",
       10, italic=True)
doc.add_paragraph()
center(f"System: {SYSTEM_TITLE}", 11, bold=True)
p("Group members:")
for i in range(1, 5):
    p(f"     {i}. ______________________________________")
p("     Date completed: ______________")
doc.add_paragraph()
p("Safety note: the live database smarttrack_db was READ-ONLY for this activity. "
  "All GRANT/REVOKE, RLS, and tests were executed in the sandbox database "
  "smarttrack_security_lab (schema copied with pg_dump --schema-only). The only "
  "cluster-wide objects created are the three deliverable NOLOGIN roles "
  "(smarttrack_super_admin, smarttrack_faculty, smarttrack_app).", italic=True)

# --------------------------------------- already-implemented vs new work --
h("Workflow status: already implemented vs. implemented by this activity", 1)
p("The controls that were ALREADY implemented in SmartTrack were left untouched "
  "(documented below as existing evidence); this activity only added the missing "
  "database-level controls and fixed one credential leak found during review.",
  italic=True)
table(
    ["Workflow item", "Status", "Where / evidence"],
    [
        ["1. Identify project roles (application RBAC)",
         "ALREADY IMPLEMENTED — unchanged",
         "backend/models.py: users.role CheckConstraint IN ('super_admin','faculty'); "
         "backend/security.py: get_current_admin()/get_current_faculty() guards on 20+ "
         "admin endpoints (users, colleges, analytics, audit, ai routes); JWT access/refresh "
         "on every protected endpoint"],
        ["7. Store credentials safely (backend secrets)",
         "ALREADY IMPLEMENTED — unchanged (evidence in Part E)",
         "backend/config.py (os.getenv + pydantic-settings env_file = \".env\"); .gitignore "
         "lines 2–3 exclude .env and backend/.env; backend/.env.example ships CHANGE_ME "
         "placeholders; user passwords hashed with bcrypt gensalt(12) (backend/security.py)"],
        ["8. Test allowed and blocked actions (application level)",
         "ALREADY IMPLEMENTED — re-executed this activity as evidence",
         "backend/tests/test_phase9_security.py (26 tests) + "
         "backend/tests/test_superadmin_guards.py (8 tests) → 31 passed, 2 skipped "
         "(missing fixture), 1 state-dependent assertion (explained in Part C); "
         "evidence: database/security/evidence/app_level_security_test_results.txt"],
        ["2. List sensitive tables (database level)",
         "NEW — this activity", "Parts B and B.1 (5 tables + access-control matrix)"],
        ["3. Create database roles",
         "NEW — this activity",
         "database/security/01_roles_and_grants.sql (smarttrack_super_admin, "
         "smarttrack_faculty, smarttrack_app — all NOLOGIN)"],
        ["4./5. Grant only needed privileges + apply schema privileges",
         "NEW — this activity",
         "01_roles_and_grants.sql: 7 GRANT + 5 REVOKE statements incl. "
         "GRANT USAGE ON SCHEMA system_core; proven by tests T-01…T-06"],
        ["6. Design one RLS policy",
         "NEW — this activity",
         "database/security/02_rls_policy.sql (travels_faculty_own_rows, "
         "travels_admin_all); proven by tests T-07…T-12"],
        ["Credential leak found during the Part E review (startup log printed the "
         "full DATABASE_URL including password)",
         "FIXED this activity",
         "backend/main.py — connection string now logged with the password masked "
         "(postgresql://postgres:****@127.0.0.1:5432/smarttrack_db)"],
    ])

# ----------------------------------------------------------------- Part A --
h("Part A — Identify Your System Roles", 1)
p("Role rule: the rows below match the application roles the backend actually "
  "recognizes (backend/models.py — CheckConstraint \"role IN ('super_admin','faculty')\", "
  "enforced by get_current_admin()/get_current_faculty() in backend/security.py). "
  "These application roles were ALREADY IMPLEMENTED before this activity — this "
  "worksheet only maps them to the database roles created in Part C. No extra roles "
  "were invented to fill the worksheet, so only two rows are filled.",
  italic=True)
table(
    ["System role", "Main responsibilities",
     "Data/actions that require access", "Security concern"],
    [
        ["super_admin (OIGE administrator)",
         "Manages user accounts, verifies/rejects uploaded documents, drives the "
         "travel workflow (approve/return/remarks), monitors compliance, reads the audit trail",
         "Full CRUD on users, travels, documents, notifications; read + append on audit_logs",
         "Greatest database access: must not hard-delete accounts (foreign keys + "
         "history) and must not rewrite audit history"],
        ["faculty (accountable traveler)",
         "Registers own mobility records, uploads pre-/post-travel documents, "
         "updates own profile, reads own notifications",
         "Own rows in travels, documents, notifications",
         "Must only ever see/edit own records; no access to other users' accounts, "
         "password hashes, or the audit trail"],
        ["— (not used)", "The backend recognizes only the two roles above.",
         "—", "Worksheet row intentionally left blank per the role rule."],
    ])
p("Database role mapping (Part C): super_admin → smarttrack_super_admin (NOLOGIN); "
  "faculty → smarttrack_faculty (NOLOGIN); backend service account → smarttrack_app (NOLOGIN, "
  "password set out-of-band by the DBA and stored only in backend/.env).")

# ----------------------------------------------------------------- Part B --
h("Part B — Select 3 to 5 Important Project Tables", 1)
p("Exact PostgreSQL names from the live schema (schema system_core); 5 tables selected.",
  italic=True)
table(
    ["Table name", "Purpose / important data", "Roles that use it",
     "Why access must be controlled"],
    [
        ["system_core.users",
         "Accounts, bcrypt password hashes, role flag, is_active, college/department",
         "super_admin (manage); faculty (own profile via backend)",
         "Credential store — password hashes and role flags must never be "
         "readable/writable by the wrong role"],
        ["system_core.travels",
         "Core mobility records (event, dates, funding, workflow status, remarks); "
         "row owner user_id",
         "super_admin (all rows); faculty (own rows)",
         "Main transactional data of the system; personal/professional travel "
         "details of faculty"],
        ["system_core.documents",
         "Uploaded pre-/post-travel evidence files, verification status, rejection remarks",
         "super_admin (verify/reject); faculty (upload own)",
         "Uploaded documents are sensitive; verification status drives compliance decisions"],
        ["system_core.notifications",
         "Per-user messages (user_id, is_read)",
         "super_admin (system-generated); faculty (own)",
         "Row-owned data — faculty must not read other users' notifications"],
        ["system_core.audit_logs",
         "Security/administrative activity history (who did what, IP, user agent)",
         "super_admin (read + append only)",
         "Audit trail: must be append-only and hidden from non-admins, or tamper "
         "evidence is lost"],
    ])

# --------------------------------------------------------------- Part B.1 --
h("Part B.1 — Access-Control Matrix", 1)
p("Y = granted, N = not granted, \"Own rows\" = table privilege plus row restriction "
  "(RLS for travels; backend row filtering for documents/notifications).", italic=True)
table(
    ["Table", "Role", "SELECT", "INSERT", "UPDATE", "DELETE", "Restriction / reason"],
    [
        ["users", "super_admin", "Y", "Y", "Y", "N",
         "Accounts are deactivated (is_active = false), never hard-deleted (FK integrity + history)"],
        ["users", "faculty", "N", "N", "N", "N",
         "No direct access at all; own profile is served by the backend"],
        ["travels", "super_admin", "Y", "Y", "Y", "Y", "Full workflow management"],
        ["travels", "faculty", "Own rows", "Own rows", "Own rows", "N",
         "RLS policy travels_faculty_own_rows; DELETE revoked (audit-trail integrity)"],
        ["documents", "super_admin", "Y", "Y", "Y", "Y", "Verification / rejection actions"],
        ["documents", "faculty", "Y", "Y", "Y", "N",
         "Backend restricts to own travels (travels.user_id join); no delete of evidence"],
        ["notifications", "super_admin", "Y", "Y", "Y", "Y", "System-generated notices"],
        ["notifications", "faculty", "Own rows", "Own rows", "Own rows (is_read)", "N",
         "Own user_id only (enforced in backend; future RLS candidate)"],
        ["audit_logs", "super_admin", "Y", "Y", "N", "N",
         "Append-only: no UPDATE/DELETE for anyone, including admin"],
        ["audit_logs", "faculty", "N", "N", "N", "N",
         "Audit trail invisible to non-admins"],
    ])

# ----------------------------------------------------------------- Part C --
h("Part C — Write and Test GRANT and REVOKE Statements", 1)
p("Full project script: database/security/01_roles_and_grants.sql (idempotent). "
  "Key statements below (at least three GRANTs and one REVOKE required):")
sql("""-- Roles (NOLOGIN group roles, 1:1 with the application roles)
CREATE ROLE smarttrack_super_admin NOLOGIN;
CREATE ROLE smarttrack_faculty NOLOGIN;

-- Schema privileges (Step 5 of the workflow)
GRANT USAGE ON SCHEMA system_core TO smarttrack_super_admin, smarttrack_faculty;

-- GRANT 1: super_admin CRUD on the five selected tables (minus the carve-outs below)
GRANT SELECT, INSERT, UPDATE, DELETE
    ON system_core.users, system_core.travels, system_core.documents,
       system_core.notifications, system_core.audit_logs
    TO smarttrack_super_admin;

-- GRANT 2: faculty read access only where needed
GRANT SELECT
    ON system_core.travels, system_core.documents, system_core.notifications
    TO smarttrack_faculty;

-- GRANT 3: faculty create/modify on its working tables
GRANT INSERT, UPDATE
    ON system_core.travels, system_core.documents, system_core.notifications
    TO smarttrack_faculty;

-- GRANT 4: sequences so faculty can allocate keys for rows it inserts
GRANT USAGE, SELECT ON ALL SEQUENCES IN SCHEMA system_core
    TO smarttrack_super_admin, smarttrack_faculty;

-- REVOKE 1: faculty can never remove mobility records
REVOKE DELETE ON system_core.travels       FROM smarttrack_faculty;
REVOKE DELETE ON system_core.documents     FROM smarttrack_faculty;
REVOKE DELETE ON system_core.notifications FROM smarttrack_faculty;

-- REVOKE 2: faculty has NO access to the user/credential table
REVOKE ALL ON system_core.users FROM smarttrack_faculty;

-- REVOKE 3: faculty cannot read or write the audit trail
REVOKE ALL ON system_core.audit_logs FROM smarttrack_faculty;

-- REVOKE 4: admin never hard-deletes accounts (deactivation instead)
REVOKE DELETE ON system_core.users FROM smarttrack_super_admin;

-- REVOKE 5: audit trail is append-only even for the admin
REVOKE UPDATE, DELETE ON system_core.audit_logs FROM smarttrack_super_admin;""")
p("Test evidence — executed in the sandbox (PostgreSQL 18.6) with "
  "database/security/run_lab_tests.ps1; full output: "
  "database/security/evidence/privilege_test_results.txt. Every test runs in its own "
  "transaction and is rolled back.", bold=True)
table(
    ["SQL / test ID", "Role + table", "Privilege change or test",
     "Actual result / interpretation", "Evidence ID"],
    [
        ["C-1", "smarttrack_super_admin + 5 selected tables",
         "GRANT SELECT/INSERT/UPDATE/DELETE, then REVOKE DELETE on users and "
         "UPDATE/DELETE on audit_logs",
         "Grant listing shows users = INSERT, SELECT, UPDATE (no DELETE) and "
         "audit_logs = INSERT, SELECT — matches the matrix",
         "EV-C1 (evidence file, grants table)"],
        ["C-2", "smarttrack_faculty + travels, documents, notifications",
         "GRANT SELECT/INSERT/UPDATE; REVOKE DELETE",
         "Listing shows INSERT, SELECT, UPDATE only — no DELETE anywhere for faculty; "
         "no rows for users/audit_logs (zero privileges)",
         "EV-C2 (evidence file, grants table)"],
        ["T-01", "super_admin + travels", "SELECT (allowed)",
         "count = 3 — admin sees all seeded rows (ALLOWED as designed)", "EV-C3"],
        ["T-02", "super_admin + travels", "DELETE disposable row 9103 (allowed)",
         "DELETE 1; remaining_disposable_rows = 0 (ALLOWED as designed)", "EV-C4"],
        ["T-03", "faculty + travels", "DELETE (blocked)",
         "ERROR: permission denied for table travels (REVOKE DELETE proven)", "EV-C5"],
        ["T-04", "faculty + users", "UPDATE is_active (blocked)",
         "ERROR: permission denied for table users (REVOKE ALL proven)", "EV-C6"],
        ["T-05", "faculty + users", "SELECT password_hash (blocked)",
         "ERROR: permission denied for table users — credential hashes unreadable", "EV-C7"],
        ["T-06", "faculty + audit_logs", "INSERT (blocked)",
         "ERROR: permission denied for table audit_logs — audit trail untouchable", "EV-C8"],
    ])
p("All expected BLOCKED errors and expected ALLOWED results occurred exactly as designed.",
  italic=True)

h("Part C (continued) — Application-level allowed/blocked tests ALREADY IMPLEMENTED "
  "before this activity", 2)
p("SmartTrack already ships two live-API security suites. They were re-executed during "
  "this activity so the evidence below is fresh; neither suite was modified.")
table(
    ["Suite (already implemented)", "What it proves", "Result in this run"],
    [
        ["backend/tests/test_phase9_security.py — 26 tests",
         "BLOCKED without auth (401): /me, /travels, /documents, /audit, /ai/insights, "
         "/compliance — plus bad token and empty bearer rejected. BLOCKED for faculty "
         "(403): read audit_logs, list users, read all remarks, create user, create "
         "announcement, create doc type, recompute/list insights, list compliance records. "
         "Input validation: 404s, bad filters, and the login SQL-injection payload test.",
         "24 passed; 2 skipped — the two ownership tests need an existing travel owned by "
         "another faculty account as fixture, and none exists in the current dev data "
         "(pre-existing condition, reported honestly)"],
        ["backend/tests/test_superadmin_guards.py — 8 tests",
         "Admin self-protection: cannot deactivate own account, cannot change own role, "
         "invalid role rejected, super-admin cap respected, deactivated users logged out "
         "immediately, updates cannot change role/active flags.",
         "7 passed; 1 state-dependent assertion (explained below)"],
        ["Run summary (both suites)",
         "Evidence file: database/security/evidence/app_level_security_test_results.txt",
         "Ran 34 tests → 31 passed, 2 skipped, 1 failed assertion. The single failure "
         "(test_role_change_blocks_account) is a data-state-dependent expectation in the "
         "PRE-EXISTING test, not a security-control failure: with only one active "
         "super_admin in the dev data, re-activating the promoted account correctly "
         "returned 200 (cap = 2 active super admins), while the test — written for a "
         "database where two super admins are already active — expected 400. All "
         "deny/guard assertions before that line passed."],
    ])

# ----------------------------------------------------------------- Part D --
h("Part D — Decide Whether Your Project Needs Row-Level Security", 1)
table(
    ["Question", "Project answer"],
    [
        ["Do multiple users/roles access the same table?",
         "Yes. Both super_admin and faculty query system_core.travels (admin dashboard "
         "vs. the \"My mobility records\" feature)."],
        ["Should a user see only rows that belong to that user, office, group, or tenant?",
         "Yes for faculty. Row ownership = travels.user_id → users.user_id (model comment: "
         "\"Record owner (faculty/admin who registered the mobility)\")."],
        ["Which table(s) would need row filtering?",
         "system_core.travels (primary); notifications and documents are secondary "
         "candidates (ownership via user_id / parent travel)."],
        ["Does the project need PostgreSQL RLS?",
         "YES"],
    ])
p("Justification based on an actual system feature: the Faculty Mobility List / "
  "\"My records\" feature must show a faculty member only their own travels while the "
  "admin sees everything. Today that filtering exists only in application code (routes add "
  "WHERE user_id = current_user.user_id), so a single missed filter or altered query clause "
  "would leak every faculty member's travel data. RLS pushes the ownership rule into the "
  "database so it holds regardless of which query runs, under the identity the backend "
  "supplies per request — and it is fail-closed when no identity is supplied.")
p("Policy design (one policy, project-specific — database/security/02_rls_policy.sql):", bold=True)
table(
    ["Design element", "Project decision"],
    [
        ["Protected table", "system_core.travels"],
        ["Row-ownership condition",
         "user_id = NULLIF(current_setting('app.current_user_id', true), '')::integer"],
        ["Permitted operations",
         "FOR ALL (SELECT/INSERT/UPDATE) limited to own rows for smarttrack_faculty; "
         "full access for smarttrack_super_admin"],
        ["How the backend supplies the identity",
         "SET LOCAL app.current_user_id = <user_id from the verified JWT> inside the "
         "request's transaction; the value comes from the authenticated token (server-side), "
         "never from SQL text supplied by the client"],
    ])
sql("""ALTER TABLE system_core.travels ENABLE ROW LEVEL SECURITY;

CREATE POLICY travels_admin_all ON system_core.travels
    FOR ALL TO smarttrack_super_admin
    USING (true) WITH CHECK (true);

CREATE POLICY travels_faculty_own_rows ON system_core.travels
    FOR ALL TO smarttrack_faculty
    USING (user_id = NULLIF(current_setting('app.current_user_id', true), '')::integer)
    WITH CHECK (user_id = NULLIF(current_setting('app.current_user_id', true), '')::integer);""")
p("Policy test evidence (sandbox):", bold=True)
table(
    ["Test", "Expectation", "Actual result", "Evidence ID"],
    [
        ["T-07 faculty SELECT with identity 9001", "allowed, own row only",
         "count = 1 (admin row invisible)", "EV-D1"],
        ["T-08 faculty SELECT with no identity", "fail-closed, 0 rows",
         "count = 0", "EV-D2"],
        ["T-09 faculty INSERT row owned by itself", "allowed", "INSERT 0 1", "EV-D3"],
        ["T-10 faculty INSERT row owned by another user", "blocked",
         "ERROR: new row violates row-level security policy for table \"travels\"", "EV-D4"],
        ["T-11 faculty UPDATE own row", "allowed", "UPDATE 1", "EV-D5"],
        ["T-12 faculty UPDATE another user's row", "blocked by RLS",
         "UPDATE 0; target row invisible", "EV-D6"],
    ])
p("Rollout note: the live backend connects as the table owner (postgres), and a table owner "
  "bypasses RLS unless FORCE ROW LEVEL SECURITY is set — so the policy was designed, executed, "
  "and proven in the sandbox first. Live enablement is a controlled follow-up: (1) backend sets "
  "SET LOCAL app.current_user_id on every request, (2) run 02_rls_policy.sql on smarttrack_db, "
  "(3) only then consider FORCE ROW LEVEL SECURITY and connecting as smarttrack_app.", italic=True)

# ----------------------------------------------------------------- Part E --
h("Part E — Protect Database Credentials in the Backend", 1)
table(
    ["Security check", "Required evidence", "Group evidence"],
    [
        ["Credentials are not hard-coded in committed source files.",
         "Screenshot or code extract showing environment-variable/config loading "
         "with secret values hidden.",
         "ALREADY IMPLEMENTED (pre-activity, unchanged): backend/config.py loads everything "
         "through os.getenv(...) / pydantic-settings with env_file = \".env\":\n"
         "    DATABASE_URL = os.getenv(\"DATABASE_URL\", \"postgresql://postgres@127.0.0.1:5432/smarttrack_db\")\n"
         "    SECRET_KEY   = os.getenv(\"SECRET_KEY\", \"your-secret-key-change-in-production\")\n"
         "Real values live only in backend/.env — actual secrets not reproduced anywhere "
         "in this report."],
        [".env or local secret files are excluded from Git.",
         "Screenshot of the relevant .gitignore entry or equivalent secret-management rule.",
         "ALREADY IMPLEMENTED (pre-activity, unchanged): .gitignore lines 2–3 exclude "
         "`.env` and `backend/.env` (plus backend/logs/ and logs/, lines 19–23, so runtime "
         "logs are never distributed). backend/.env.example ships only CHANGE_ME "
         "placeholders. User passwords are also already protected at rest with bcrypt "
         "(gensalt(12), backend/security.py) and the password_hash column is never "
         "returned by the API — DB-level denial of it is proven by test T-05."],
        ["Repository and screenshots expose no password, token, or private connection string.",
         "Repository review and sanitized evidence.",
         "Repository-wide review found no literal password/token/private connection string. "
         "One finding was found and fixed during this activity: the startup log line "
         "logger.info(f\"Database: {settings.DATABASE_URL}\") printed the full connection "
         "string with password — it now masks it: postgresql://postgres:****@127.0.0.1:5432/"
         "smarttrack_db (backend/main.py)."],
        ["Database account has only the privileges needed by the backend.",
         "Short explanation linked to the access-control matrix.",
         "Role smarttrack_app (01_roles_and_grants.sql, Step 5): SELECT/INSERT/UPDATE/DELETE "
         "on schema system_core + USAGE/SELECT on sequences, NO superuser/createrole/createdb "
         "and no schema DDL — the runtime union of the two matrix roles; it may also SET ROLE "
         "into smarttrack_super_admin/smarttrack_faculty. NEW this activity (deliverable of "
         "Part C/Part E); activation (password set out-of-band → backend/.env) is a "
         "deliberate follow-up so this activity does not disturb the running system."],
    ])
p("Summary: credential loading, Git exclusion, secret placeholders, and password hashing "
  "were ALREADY IMPLEMENTED and were left untouched; the only change made by this activity "
  "is the masked startup log line (finding above).", italic=True)
p("Scope note — test fixture (left unchanged by group decision): backend/tests/base.py "
  "carries default credentials for the group's localhost-only development test accounts used "
  "by the live-API test suites (seeded admin/faculty fixture users). These are test fixtures, "
  "not application or database credentials. All real secrets (database URL, JWT signing key) "
  "remain only in the gitignored backend/.env, and no secret value appears in this report or "
  "in the committed evidence files (verified by a secret-pattern scan: 0 hits).", italic=True)
p("Safe pattern in this project (example only — real values hidden):", bold=True)
sql("""# backend/config.py  (SAFE PATTERN - secrets come from environment)
import os
DATABASE_URL = os.getenv("DATABASE_URL", "postgresql://postgres@127.0.0.1:5432/smarttrack_db")
# Do not place a real password directly in committed source code.""")
p("Credential-safety rule observed: no actual password, token, private connection string, "
  "or .env content is submitted or displayed in this report; the evidence proves the "
  "protection method only.", italic=True)

# ----------------------------------------------------------------- Part F --
h("Part F — Review One Project Query for SQL Injection Protection", 1)
table(
    ["Review item", "Group response"],
    [
        ["Feature / file / function reviewed",
         "Login — backend/routes/auth.py, login()"],
        ["User input involved",
         "request.email (and request.password) from the untrusted JSON request body"],
        ["SQL or ORM statement used",
         "user = db.query(User).filter(User.email == request.email).first()\n"
         "→ SQLAlchemy compiles to: SELECT ... FROM system_core.users "
         "WHERE system_core.users.email = %(email)s"],
        ["Parameterized query / ORM binding used?", "YES"],
        ["Why the statement avoids SQL injection, or correction made",
         "The email value is passed to the driver as a bind parameter, separately from the "
         "SQL template: PostgreSQL receives a prepared statement plus data, and the data is "
         "never re-parsed as SQL syntax. An input like ' OR '1'='1 is just an email value "
         "that matches no row (generic \"Invalid email or password\" response), and "
         "the password is then checked in Python with bcrypt.checkpw() against the stored "
         "hash — never inside SQL. A repository-wide search found zero f-string/format-built "
         "SQL statements in project backend code; all DB access goes through the ORM with "
         "bound parameters. No correction was needed for this query."],
    ])

p("Already-implemented regression test for this exact concern (re-run this activity: PASSED): "
  "backend/tests/test_phase9_security.py → test_login_sql_injection_safe posts the email "
  "payload \"oige@csucc.edu.ph'; DROP TABLE users;--\" and asserts 401 or 422 plus a healthy "
  "server afterwards — result in this run: ok "
  "(database/security/evidence/app_level_security_test_results.txt).")

# ----------------------------------------------------- Evidence checklist --
h("Required Evidence Checklist (activity item 10)", 1)
table(
    ["Required evidence", "Where it is"],
    [
        ["System roles + 3–5 important tables, actual project names",
         "Parts A & B (from backend/models.py and the live system_core schema)"],
        ["Completed access-control matrix for selected tables and actual system roles",
         "Part B.1 (+ grant listing EV-C1/EV-C2)"],
        ["Project-specific GRANT/REVOKE SQL and execution/test result",
         "Part C; database/security/01_roles_and_grants.sql; "
         "database/security/evidence/privilege_test_results.txt"],
        ["RLS YES/NO decision with project-specific justification; policy design/testing",
         "Part D; database/security/02_rls_policy.sql (tests T-07…T-12)"],
        ["Backend credential-protection evidence with all secret values hidden",
         "Part E (backend/config.py, backend/.env.example, .gitignore, redacted main.py log line)"],
        ["One actual project query with SQL-injection prevention explanation",
         "Part F (backend/routes/auth.py login query) + already-implemented "
         "test_login_sql_injection_safe (re-run, passed)"],
        ["Application-level allowed/blocked tests (already implemented, re-executed)",
         "database/security/evidence/app_level_security_test_results.txt "
         "(34 tests: 31 passed, 2 skipped on missing fixture, 1 state-dependent assertion "
         "explained in Part C)"],
        ["GitHub checkpoint / submitted SQL-security documentation matching the demo",
         "database/security/*.sql, run_lab_tests.ps1, evidence file, and "
         "docs/IT108_ACTIVITY7_DATABASE_SECURITY.md"],
    ])
p("Reproduce the demonstration (live database stays read-only):", bold=True)
sql("""powershell -ExecutionPolicy Bypass -File database\\security\\run_lab_tests.ps1
# rebuilds sandbox smarttrack_security_lab, applies roles/GRANTs + RLS,
# seeds safe test rows, runs all 12 tests, rewrites the evidence file.""")

# ------------------------------------------------------------ reflection --
h("Post-Laboratory Reflection", 1)
reflections = [
    ("1. Which project role currently has the greatest database access, and is every "
     "privilege necessary?",
     "super_admin — via smarttrack_super_admin it holds CRUD on the five selected tables. "
     "Two privileges were deliberately removed even for this role: DELETE on users (accounts "
     "are deactivated, not deleted, to preserve foreign keys and history) and UPDATE/DELETE "
     "on audit_logs (append-only trail). Its audit_logs access itself was cut down to "
     "SELECT + INSERT only, so every remaining privilege is necessary for its tasks."),
    ("2. Which selected table contains the most sensitive or critical data, and how is "
     "access to it restricted?",
     "system_core.users — it stores bcrypt password hashes and role flags. Faculty has zero "
     "privileges on it (REVOKE ALL; proven by tests T-04 and T-05: permission denied), the "
     "admin cannot hard-delete rows (REVOKE DELETE), and hashes are never returned by the API."),
    ("3. Why did your group decide to use or not use Row-Level Security?",
     "Use (YES). Both roles read the same travels table while faculty must only see their own "
     "records; today that rule lives only in application code. RLS enforces ownership inside "
     "the database itself, is fail-closed when the backend supplies no identity (test T-08 → "
     "0 rows), and its identity model matches our authentication architecture "
     "(JWT-derived SET LOCAL app.current_user_id per request)."),
    ("4. What would happen if the backend database credentials were committed publicly, and "
     "what control prevents that in your project?",
     "Anyone with the credentials could read or destroy every mobility record and the audit "
     "trail — and with a superuser-style account, alter the schema as well. Prevention in this "
     "project: credentials live only in backend/.env, which .gitignore excludes (lines 2–3); "
     "the committed backend/.env.example carries CHANGE_ME placeholders only; the startup log "
     "line that used to print the full connection string was masked (Part E); and the "
     "least-privilege smarttrack_app role limits the blast radius of any leak."),
    ("5. Which project query did you review for SQL injection, and what specific mechanism "
     "keeps user input separate from SQL syntax?",
     "The login query db.query(User).filter(User.email == request.email) in "
     "backend/routes/auth.py. SQLAlchemy sends it as a parameterized statement "
     "(WHERE system_core.users.email = %(email)s): the database parses and executes the SQL "
     "template and compares the separately-bound value, so user input can never change the "
     "statement's syntax — a classic ' OR '1'='1 payload is only ever data."),
]
for question, answer in reflections:
    p(question, bold=True)
    p(answer)

doc.core_properties.title = "IT108 Activity 7 — PostgreSQL Database Security and Access Control"
doc.core_properties.subject = SYSTEM_TITLE
doc.save(OUT)
print(f"Written: {OUT}")
